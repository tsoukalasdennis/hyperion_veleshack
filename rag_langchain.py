from pathlib import Path

from docx import Document as DocxDocument
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from embedding_adapter import NomicEmbeddings


DOCS_DIR = Path("docs")


class RAG:
    """LangChain + FAISS document retrieval."""

    def __init__(
        self,
        docs_dir: Path = DOCS_DIR,
        chunk_size: int = 1000,
        chunk_overlap: int = 100,
    ):
        self.docs_dir = docs_dir
        self.embeddings = NomicEmbeddings()

        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
        )

        self.vector_store = self._build_index()

    def _load_documents(self) -> list[Document]:
        """Load DOCX files as LangChain Documents."""

        documents = []

        for path in sorted(self.docs_dir.glob("*.docx")):
            document = DocxDocument(path)

            paragraphs = [
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            text = "\n".join(paragraphs)

            documents.append(
                Document(
                    page_content=text,
                    metadata={"source": path.name},
                )
            )

        return documents

    def _build_index(self):
        """Load, split, embed, and index all documents."""

        documents = self._load_documents()

        print(f"Loaded {len(documents)} documents")

        chunks = self.splitter.split_documents(documents)

        print(f"Created {len(chunks)} chunks")

        vector_store = FAISS.from_documents(
            chunks,
            self.embeddings,
        )

        print("FAISS index ready")

        return vector_store

    def search(self, query: str, top_k: int = 3) -> list[Document]:
        """Return the most relevant document chunks."""

        return self.vector_store.similarity_search(
            query,
            k=top_k,
        )


if __name__ == "__main__":
    rag = RAG()

    query = "What is the HYPER-AI Resource Model (HRM)?"

    results = rag.search(query)

    print(f"\nSearching for: {query}\n")

    for number, result in enumerate(results, start=1):
        print(f"--- Result {number} ---")
        print(f"Source: {result.metadata['source']}")
        print(result.page_content[:500])
        print()
